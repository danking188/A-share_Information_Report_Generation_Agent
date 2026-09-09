"""
Word报告生成模块
将分析结果生成格式化的Word文档
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from datetime import datetime
from typing import Dict, List
import os
import platform
from loguru import logger


class ReportGenerator:
    """Word报告生成器"""

    def __init__(self):
        self.document = Document()
        if platform.system() == 'Darwin':
            default_body_font = 'Songti SC'
            default_heading_font = 'PingFang SC'
        elif platform.system() == 'Windows':
            default_body_font = 'Microsoft YaHei'
            default_heading_font = 'Microsoft YaHei'
        else:
            default_body_font = 'Noto Sans CJK SC'
            default_heading_font = 'Noto Sans CJK SC'
        self.body_font = os.getenv('REPORT_BODY_FONT') or default_body_font
        self.heading_font = os.getenv('REPORT_HEADING_FONT') or default_heading_font
        self._setup_styles()

    def _set_run_font(self, run, font_name: str, size: Pt = None):
        run.font.name = font_name
        if size is not None:
            run.font.size = size
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:ascii'), font_name)
        fonts.set(qn('w:hAnsi'), font_name)
        fonts.set(qn('w:eastAsia'), font_name)

    def _setup_styles(self):
        """设置文档样式"""
        try:
            # 设置默认字体
            self.document.styles['Normal'].font.name = self.body_font
            self.document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), self.body_font)
            self.document.styles['Normal'].font.size = Pt(10.5)

            title_style = self.document.styles['Title']
            title_style.font.name = self.heading_font
            title_style.font.color.rgb = RGBColor(0, 0, 0)
            title_style._element.rPr.rFonts.set(qn('w:eastAsia'), self.heading_font)
            title_ppr = title_style._element.get_or_add_pPr()
            title_border = title_ppr.find(qn('w:pBdr'))
            if title_border is not None:
                title_ppr.remove(title_border)

            # 修改现有标题样式（紧凑版）
            for i in range(1, 4):
                try:
                    heading_style = self.document.styles[f'Heading {i}']
                    heading_style.font.name = self.heading_font
                    heading_style._element.rPr.rFonts.set(qn('w:eastAsia'), self.heading_font)
                    heading_style.font.size = Pt(14 - i * 1)  # 更紧凑的字体大小
                    heading_style.font.bold = True
                    heading_style.font.color.rgb = RGBColor(0, 0, 0)
                    heading_style.paragraph_format.space_before = Pt(6)  # 减少间距
                    heading_style.paragraph_format.space_after = Pt(3)
                except KeyError:
                    heading_style = self.document.styles.add_style(
                        f'Heading {i}',
                        WD_STYLE_TYPE.PARAGRAPH
                    )
                    heading_style.font.name = self.heading_font
                    heading_style._element.rPr.rFonts.set(qn('w:eastAsia'), self.heading_font)
                    heading_style.font.size = Pt(14 - i * 1)
                    heading_style.font.bold = True
                    heading_style.font.color.rgb = RGBColor(0, 0, 0)
                    heading_style.paragraph_format.space_before = Pt(6)
                    heading_style.paragraph_format.space_after = Pt(3)
        except Exception as e:
            logger.warning(f"设置样式失败: {e}")
            pass

    def _add_title(self, text: str, level: int = 1):
        """添加标题"""
        heading = self.document.add_heading(text, level=level)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return heading

    def _add_paragraph(self, text: str, bold: bool = False, font_size: Pt = None, indent: bool = False):
        """添加段落"""
        paragraph = self.document.add_paragraph()
        run = paragraph.add_run(text)
        self._set_run_font(run, self.body_font, font_size or Pt(10.5))
        run.font.bold = True if bold else False
        run.font.color.rgb = RGBColor(0, 0, 0)

        # 设置段落格式
        paragraph.paragraph_format.line_spacing = 1.15  # 紧凑行距
        paragraph.paragraph_format.space_after = Pt(3)  # 减少间距

        # 仅对长段落设置首行缩进
        if indent and len(text) > 50:
            paragraph.paragraph_format.first_line_indent = Cm(0.35)

        return paragraph

    def _add_bullet_points(self, items: List[str]):
        """添加项目符号列表"""
        for item in items:
            paragraph = self.document.add_paragraph(item, style='List Bullet')
            paragraph.paragraph_format.left_indent = Cm(0.5)

    def _add_table(self, data: List[List[str]], headers: List[str] = None):
        """添加表格"""
        if not data:
            return

        rows = len(data)
        cols = len(data[0]) if data else 0

        table = self.document.add_table(rows=rows, cols=cols)
        table.style = 'Light Grid Accent 1'

        for i, row_data in enumerate(data):
            row = table.rows[i]
            for j, cell_data in enumerate(row_data):
                cell = row.cells[j]
                cell.text = str(cell_data)

                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        self._set_run_font(run, self.body_font, Pt(9))

        if headers and len(headers) == cols:
            header_row = table.rows[0]
            for i, header in enumerate(headers):
                cell = header_row.cells[i]
                cell.text = header
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.size = Pt(10)

        return table

    def _add_page_break(self):
        """添加分页符"""
        self.document.add_page_break()

    def add_report_title(self, stock_name: str, stock_code: str):
        """添加报告标题（替代封面页）"""
        # 标题
        title = self.document.add_paragraph(style='Title')
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = title.add_run(f'{stock_name}（{stock_code}）研究报告')
        self._set_run_font(title_run, self.heading_font, Pt(16))
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0, 0, 0)

        # 日期
        date_paragraph = self.document.add_paragraph()
        date_run = date_paragraph.add_run(f'报告日期：{datetime.now().strftime("%Y年%m月%d日")}')
        self._set_run_font(date_run, self.body_font, Pt(10))
        date_run.font.color.rgb = RGBColor(128, 128, 128)

        self.document.add_paragraph()  # 空行

    def generate_investment_advice_section(self, advice_text: str):
        """生成投资建议部分"""
        self._add_title('投资建议', level=1)

        # 直接显示投资建议文本（评级+核心逻辑合并）
        self._add_paragraph(advice_text, indent=True)

        self.document.add_paragraph()

    def generate_investment_logic_section(self, logic_text: str):
        """生成投资逻辑部分"""
        self._add_title('投资逻辑', level=1)

        # 显示投资逻辑文本
        self._add_paragraph(logic_text, indent=True)

        self.document.add_paragraph()

    def generate_full_report(self, report_data: Dict) -> Document:
        """生成完整报告"""
        logger.info("开始生成Word文档...")

        try:
            stock_name = report_data.get('stock_name', '未知公司')
            stock_code = report_data.get('symbol', 'N/A')

            # 1. 添加标题（替代封面页）
            self.add_report_title(stock_name, stock_code)

            quality = report_data.get('data_quality') or {}
            if quality:
                status_labels = {'passed': '完整', 'degraded': '部分缺失', 'failed': '不合格'}
                status = status_labels.get(quality.get('status'), quality.get('status', '未知'))
                warnings = quality.get('warnings') or []
                quality_text = f"数据状态：{status}（质量评分：{quality.get('score', 'N/A')}）"
                if warnings:
                    quality_text += "；" + "；".join(warnings)
                self._add_paragraph(quality_text, font_size=Pt(9))

            # 获取sections
            sections = report_data.get('sections', {})

            # 2. 投资建议（放在第一位）
            if 'investment_advice' in sections:
                self.generate_investment_advice_section(sections['investment_advice'])

            # 3. 投资逻辑（新增）
            if 'investment_logic' in sections:
                self.generate_investment_logic_section(sections['investment_logic'])

            # 4. 公司概况
            if 'company_overview' in sections:
                self._add_title('公司概况', level=1)
                self._add_paragraph(sections['company_overview'], indent=True)
                self.document.add_paragraph()

            # 4. 财务数据分析
            self._add_title('一、财务数据分析', level=1)

            if 'financial_indicators' in sections:
                self._add_title('1. 主要财务指标（近三年）', level=2)
                indicator_text = sections['financial_indicators']
                lines = indicator_text.split('\n')
                for line in lines:
                    line = line.strip()
                    if line and not line.startswith('###'):
                        self._add_paragraph(line, indent=False)

            if 'financial_analysis' in sections:
                self._add_title('2. 财务数据分析', level=2)
                self._add_paragraph(sections['financial_analysis'], indent=True)

            self.document.add_paragraph()

            # 5. 业务展望和行业地位
            if 'business_outlook' in sections:
                self._add_title('二、业务展望和行业地位', level=1)
                self._add_paragraph(sections['business_outlook'], indent=True)
                self.document.add_paragraph()

            # 6. 可比上市公司对比
            if 'comparable_analysis' in sections:
                self._add_title('三、可比上市公司对比', level=1)
                self._add_paragraph(sections['comparable_analysis'], indent=True)
                self.document.add_paragraph()

            sources = report_data.get('sources') or []
            if sources:
                self._add_title('资料来源', level=1)
                seen = set()
                for source in sources:
                    key = (source.get('evidence_id'), source.get('source'))
                    if key in seen:
                        continue
                    seen.add(key)
                    self._add_paragraph(
                        f"[{source.get('evidence_id')}] {source.get('source')} "
                        f"（片段：{source.get('chunk_id')}）",
                        font_size=Pt(9),
                    )

            logger.info("Word文档生成完成！")
            return self.document

        except Exception as e:
            logger.error(f"生成报告失败: {e}")
            raise

    def save_report(self, output_path: str):
        """保存报告到文件"""
        try:
            self.document.save(output_path)
            logger.info(f"报告已保存至: {output_path}")
        except Exception as e:
            logger.error(f"保存报告失败: {e}")
            raise


if __name__ == "__main__":
    # 测试代码
    generator = ReportGenerator()

    test_data = {
        'stock_name': '凯文教育',
        'symbol': '002659',
        'sections': {
            'investment_advice': '投资建议：推荐。公司业绩良好，行业地位稳固，具有增长潜力。',
            'company_overview': '这是一家教育服务公司。',
            'financial_indicators': '### 主要财务指标\n营业收入：...',
            'financial_analysis': '财务状况良好...',
            'business_outlook': '业务前景广阔...',
            'comparable_analysis': '相比同行业公司...'
        }
    }

    doc = generator.generate_full_report(test_data)

    output_path = os.path.join(os.path.dirname(__file__), 'output', 'test_report.docx')
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    generator.save_report(output_path)

    print(f"测试报告已生成: {output_path}")
